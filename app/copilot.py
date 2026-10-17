"""Murmur Copilot: live interview assistant.

While a copilot session records, the interviewer's questions (computer audio) are detected and an
answer is written for you in real time, from your resume, your documents, the job description and
your answer preferences. Resumes and documents are kept locally in data/copilot/.
"""

import html
import json
import logging
import re
import threading
import time
import urllib.request
import uuid
from datetime import datetime
from pathlib import Path

from .config import DATA_DIR
from . import summary

log = logging.getLogger("murmur.copilot")

LIBRARY_DIR = DATA_DIR / "copilot"
KINDS = ("resumes", "documents")
_lock = threading.RLock()

CLAUDE_MODELS = ["claude-haiku-5-5", "claude-sonnet-5-5", "claude-opus-5-5"]

# Your own default instructions (you can edit them per session; the last ones you used are kept).
DEFAULT_INSTRUCTIONS = """You are my interview assistant. Write the answers from my perspective, in the first person, in simple and natural language.

Answer the question directly first. Then add the reasoning and, if useful, one concrete example.

Usually write 3-5 short sentences. One main idea per sentence. Avoid long sentences with many commas.

For experience questions use this order: What was the problem? What did I do? Why? What was the result? Focus on my concrete actions.

For technical questions, first explain the term in plain words. Then describe how it works or give a practical example. Do not add details nobody asked for.

Use only confirmed facts about my experience. Never invent numbers, results or implementations. Separate work I actually did from solutions I would propose.

Present me as a senior engineer who builds things and takes ownership. Show it through concrete decisions and tasks, not through generic self-praise.

The answers should sound like a natural conversation. Avoid textbook language, unnecessary intros and repetition. Only go deeper when the question asks for it.

Write only the answer I can say out loud. No extra comments, variants or phrases like "You could say". Stop as soon as the question is answered."""

DEFAULT_PREFS = {"format": "full", "length": "balanced", "tone": "simple", "star": True, "fillers": True}


# ---------------------------------------------------------------- library: resumes and documents

def _dir(kind: str) -> Path:
    if kind not in KINDS:
        raise ValueError("Unknown library")
    d = LIBRARY_DIR / kind
    d.mkdir(parents=True, exist_ok=True)
    return d


def _save(kind: str, item: dict) -> dict:
    with _lock:
        path = _dir(kind) / f"{item['id']}.json"
        tmp = path.with_suffix(".json.tmp")
        tmp.write_text(json.dumps(item, ensure_ascii=False, indent=1), encoding="utf-8")
        tmp.replace(path)
    return item


def get_item(kind: str, item_id: str) -> dict | None:
    path = _dir(kind) / f"{re.sub(r'[^a-z0-9]', '', item_id)}.json"
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def list_items(kind: str) -> list[dict]:
    items = []
    for path in _dir(kind).glob("*.json"):
        try:
            it = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        items.append({k: it.get(k) for k in ("id", "name", "source", "created_at", "filename", "size")}
                     | {"preview": it.get("text", "")[:600], "chars": len(it.get("text", ""))})
    return sorted(items, key=lambda x: x["created_at"] or "", reverse=True)


def add_item(kind: str, name: str, text: str, source: str, filename: str = "", size: int = 0) -> dict:
    text = text.strip()
    if not text:
        raise ValueError("No text found")
    item = {"id": uuid.uuid4().hex[:12], "name": name.strip() or "Untitled", "source": source, "text": text,
            "filename": filename, "size": size, "created_at": datetime.now().isoformat(timespec="seconds")}
    return _save(kind, item)


def update_item(kind: str, item_id: str, name: str | None, text: str | None) -> dict:
    item = get_item(kind, item_id)
    if not item:
        raise KeyError(item_id)
    if name is not None and name.strip():
        item["name"] = name.strip()
    if text is not None:
        item["text"] = text.strip()
    return _save(kind, item)


def delete_item(kind: str, item_id: str) -> None:
    with _lock:
        for path in _dir(kind).glob(f"{re.sub(r'[^a-z0-9]', '', item_id)}.*"):
            path.unlink(missing_ok=True)


def extract_text(path: Path) -> str:
    """Text from an uploaded PDF, Word document or plain text file."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        return "\n\n".join((page.extract_text() or "").strip() for page in PdfReader(str(path)).pages)
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    if ext in (".txt", ".md", ".markdown", ".csv", ".json", ".tex", ".rtf", ".html", ".htm"):
        raw = path.read_text(encoding="utf-8", errors="replace")
        return html_to_text(raw) if ext in (".html", ".htm") else raw
    raise ValueError("Use a PDF, Word (.docx) or text file")


def html_to_text(raw: str) -> str:
    raw = re.sub(r"(?is)<(script|style|noscript|svg|head)[^>]*>.*?</\1>", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>|</(p|div|li|h\d|tr|section|article)>", "\n", raw)
    text = html.unescape(re.sub(r"<[^>]+>", " ", raw))
    lines = [re.sub(r"[ \t\r\f\v]+", " ", line).strip() for line in text.split("\n")]
    return "\n".join(line for line in lines if line)


def fetch_url(url: str) -> tuple[str, str]:
    """(page title, page text) of a web page, e.g. a job post or a company page."""
    if not re.match(r"^https?://", url):
        url = "https://" + url
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Murmur Copilot"})
    with urllib.request.urlopen(req, timeout=20) as r:
        raw = r.read(3_000_000).decode(r.headers.get_content_charset() or "utf-8", errors="replace")
    m = re.search(r"(?is)<title[^>]*>(.*?)</title>", raw)
    return (html.unescape(m.group(1)).strip() if m else url), html_to_text(raw)


# ---------------------------------------------------------------- language models

def pick_provider(settings: dict, provider: str | None) -> str:
    ready = summary.available_providers(settings)
    if not ready:
        raise RuntimeError("No API key yet. Add one in Settings, API keys")
    return provider if provider in ready else (settings.get("summary_provider") if settings.get("summary_provider") in ready else ready[0])


_no_effort: set[str] = set()   # models that reject reasoning_effort


def _openai_content(content):
    if isinstance(content, str):
        return content
    out = []
    for part in content:
        if part["type"] == "text":
            out.append({"type": "text", "text": part["text"]})
        else:
            out.append({"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{part['data']}", "detail": "high"}})
    return out


def _claude_content(content):
    if isinstance(content, str):
        return content
    out = []
    for part in content:
        if part["type"] == "text":
            out.append({"type": "text", "text": part["text"]})
        else:
            out.append({"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": part["data"]}})
    return out


def stream_chat(settings: dict, provider: str | None, model: str, system: str, messages: list[dict], on_delta,
                fast: bool = False, cancelled=lambda: False) -> str:
    """Stream a reply; on_delta(text) for every piece. messages: [{"role", "content": str | parts}], where a
    part is {"type": "text", "text"} or {"type": "image", "data": base64 jpeg}. fast=True asks reasoning
    models not to think first (first word in ~1 s instead of ~3 s)."""
    provider = pick_provider(settings, provider)
    out = []
    if provider == "claude":
        import anthropic

        ws = settings.get("anthropic_workspace_id", "")
        client = anthropic.Anthropic(api_key=settings["anthropic_api_key"],
                                     default_headers={"anthropic-workspace-id": ws} if ws else None)
        msgs = [{"role": m["role"], "content": _claude_content(m["content"])} for m in messages]
        with client.messages.stream(model=model if model in CLAUDE_MODELS else CLAUDE_MODELS[0], max_tokens=2500,
                                    system=system, messages=msgs) as stream:
            for piece in stream.text_stream:
                if cancelled():
                    break
                out.append(piece)
                on_delta(piece)
        return "".join(out)

    import openai

    model = model or settings.get(f"{provider}_model", "")
    if not model:
        raise RuntimeError("Choose a model for this session")
    client = summary._client(provider, settings[f"{provider}_api_key"]).with_options(timeout=120)
    msgs = [{"role": "system", "content": system}] + [{"role": m["role"], "content": _openai_content(m["content"])} for m in messages]
    extra = {}
    if provider == "openai" and model not in _no_effort:
        extra["reasoning_effort"] = "none" if fast else "low"
    try:
        stream = client.chat.completions.create(model=model, stream=True, messages=msgs, **extra)
    except openai.BadRequestError as exc:
        if not extra or "reasoning" not in str(exc).lower():
            raise
        _no_effort.add(model)            # older model: no reasoning setting
        stream = client.chat.completions.create(model=model, stream=True, messages=msgs)
    try:
        for chunk in stream:
            if cancelled():
                break
            if chunk.choices and chunk.choices[0].delta and chunk.choices[0].delta.content:
                piece = chunk.choices[0].delta.content
                out.append(piece)
                on_delta(piece)
    finally:
        stream.close()
    return "".join(out)


def complete(settings: dict, system: str, user: str, provider: str | None = None, model: str = "") -> str:
    return stream_chat(settings, provider, model, system, [{"role": "user", "content": user}], lambda _: None, fast=True)


def job_from_page(settings: dict, url: str) -> dict:
    """Company, job title and description from a job post link."""
    title, text = fetch_url(url)
    text = text[:30000]
    try:
        raw = complete(settings, "Extract the job post from this web page text. Reply with only JSON: "
                       '{"company": "...", "title": "...", "description": "..."}. The description is the full '
                       "job description (tasks, requirements, benefits) as plain text, in the page's language.",
                       f"Page title: {title}\n\n{text}")
        data = summary._extract_json(raw)
        return {"company": str(data.get("company", "")).strip(), "title": str(data.get("title", "")).strip(),
                "description": str(data.get("description", "")).strip() or text[:8000]}
    except Exception:
        log.exception("Job extraction by the model failed; using the page text")
        return {"company": "", "title": title, "description": text[:8000]}


# ---------------------------------------------------------------- answers

LANG_NAMES = {"en": "English", "de": "German", "az": "Azerbaijani"}

# How answers must sound. This is the "training": a precise style guide plus examples, sent with every
# request (identical prefix, so the provider caches it and it costs almost no extra time).
STYLE_GUIDE = """You are Murmur Copilot. During a live {kind} you write what the candidate says next. The candidate \
reads your text from the screen while talking, so it must sound like a person speaking, not like a text being \
read out.

HOW IT MUST SOUND
- Natural and spoken, like a confident engineer talking to a colleague. Never academic, never textbook, never \
like a prepared speech. No "Great question", no "In conclusion", no summary at the end.
- Simple words and short sentences, about B1-B2 level, so a non-native speaker can say it easily and the \
listener follows. One idea per sentence. No long sentences full of commas, no rare words, no idioms, no buzzwords.
- First person: "I", and "we" for the team.
- No headings, no labels like "Situation:" or "Task:", no bold text.

FIRST LINE
- Start your output with one line "# " plus a short title of the question in the answer language (3-8 words, \
e.g. "# Conflict with a colleague"), then a new line, then the answer. The title shows the candidate which \
question you understood.

STRUCTURE
- The first one or two sentences answer the question directly, so the candidate can start talking at once.
- Then the substance: why, how it works, or the story.
- Match the length to the question. Greetings, small talk and simple yes/no questions get one or two sentences; \
do not stretch them. A normal technical or experience question gets a full, concrete answer of about 90-140 \
words, a big one (design, a whole project) up to about 170. Short paragraphs of two to four sentences.
- Speech recognition often garbles names ("Kicilluak" = Keycloak, "Cuba Nettis" = Kubernetes, "Jason" = JSON). \
When the likely meaning is clear, just answer it as if heard correctly. Only if the question is really unclear \
or cut off, write one short polite sentence asking them to repeat or rephrase it (in the answer language).

CONCRETE EXAMPLES AND STORIES
- Back every claim with a concrete situation from the candidate's real experience in the resume and documents: \
which project, what the problem was, what exactly I did and why, and what came out of it. Tell it like a short \
story: "In project X we had problem Y. I did Z, because ... In the end ..."
- Behavioural questions: one story in this order: what was going on, what I did, the result and what I learned. \
That is the STAR order, but never say the labels and do not force it on questions that are not about experience. \
Build the story only from situations that the resume and documents really describe. If none fits exactly (for \
example "a mistake you made"), take the closest real project and tell what typically went wrong or was hard \
there and how I handle it, without inventing a specific incident, date, number or outcome. Never comment on \
your own answer or its sources ("I can't point to...", "the documents don't say...", "I don't have a figure", "dazu kann ich keine konkrete Implementierung beschreiben") in any language; \
just tell the closest real situation naturally, the way a person would.
- Technical questions: first explain it simply in one or two sentences, then show where I used it in practice \
("At my last company I used this when ..."). If no project of mine fits, give a realistic example of how it is used, \
without claiming I did it.
- Good storytelling: a clear problem, what I decided, a concrete result. Real details (technology, team, what \
changed) make it believable. Never invent numbers, employers, projects or results that are not in the documents. \
When the documents give no number, describe the result in words ("the export became stable", "the team stopped \
losing time on it") and never say that you have no figure, no data or nothing "to claim": a real candidate \
would not say that.

SCREENSHOTS
- A screenshot with a coding task: one or two sentences about the approach in plain words, then the code in a \
fenced code block, then one sentence on time and space complexity.
- A screenshot with something else (a question, a diagram, a document): answer about what it shows, in the same \
spoken style. If earlier screenshots are attached, use them as context (e.g. a task that continues).

STYLE EXAMPLES (only to show the tone; they are not the candidate's facts)
Question: "Tell me about a conflict in your team."
Bad: "Situation: In my previous role, a disagreement arose regarding architectural paradigms. Task: ..."
Good: "Sure. In one project we couldn't agree on the search setup. A colleague wanted a separate database, I \
thought our Postgres was enough. So I built a small test with real data over two days. It showed Postgres was \
fast enough, and we kept the system simpler. What I took from it: test it, don't argue about it."
Question: "Do you know Docker?"
Good: "Yes, I use it every day. All my services run in Docker, locally and in the pipeline."
"""

FORMAT_RULES = {
    "full": "Write the answer as full spoken sentences the candidate can say as they are.",
    "script_bullets": "Write one or two spoken opening sentences, then 3-5 short bullet points the candidate "
                      "can talk about in their own words.",
    "bullets": "After the first spoken sentence, use only short bullet points (keywords, short phrases).",
}
LENGTH_RULES = {
    "short": "Keep every answer short: at most about 60 words, even for hard questions.",
    "balanced": "Follow the lengths above (about 90-140 words for a normal question).",
    "long": "For real questions you may go to about 180-260 words with a full story or explanation; easy "
            "questions stay short.",
}
TONE_RULES = {"simple": "Tone: relaxed and friendly, plain words.",
              "formal": "Tone: polite and professional, still simple and spoken, never stiff."}


def build_system(cfg: dict) -> str:
    prefs = {**DEFAULT_PREFS, **(cfg.get("prefs") or {})}
    lang = cfg.get("language", "auto")
    parts = [
        STYLE_GUIDE.format(kind="job interview" if cfg.get("type", "interview") == "interview" else "call"),
        "THIS SESSION",
        "The conversation comes from live speech recognition: 'Interviewer' is the other side, 'You' is the "
        "candidate. Recognition errors happen; infer what was meant.",
        FORMAT_RULES.get(prefs["format"], FORMAT_RULES["full"]),
        LENGTH_RULES.get(prefs["length"], LENGTH_RULES["balanced"]),
        TONE_RULES.get(prefs["tone"], TONE_RULES["simple"]),
        "A filler word now and then (\"well\", \"so\", \"honestly\") is fine, at most one per answer."
        if prefs.get("fillers") else "No filler words.",
        "" if prefs.get("star", True) else "Do not use the STAR order; answer experience questions more briefly.",
        (f"LANGUAGE: answer in {LANG_NAMES[lang]}." if lang in LANG_NAMES
         else "LANGUAGE: answer in the language the interviewer used for the latest question (English question, "
              "English answer; German question, German answer).")
        + " The whole answer stays in that one language, and this rule wins over any language named in the "
          "candidate's instructions below (their other style rules still apply).",
        "Output only what the candidate says (Markdown only for code blocks or bullet points).",
    ]
    if cfg.get("instructions", "").strip():
        parts.append("The candidate's own instructions (follow them; where they conflict with the rules above, "
                     "they win):\n<instructions>\n" + cfg["instructions"].strip() + "\n</instructions>")
    job = [f"Company: {cfg.get('company', '')}", f"Role: {cfg.get('title', '')}"]
    if cfg.get("job_description", "").strip():
        job.append("Job description:\n" + cfg["job_description"].strip())
    parts.append("<job>\n" + "\n".join(job) + "\n</job>")
    if cfg.get("_resume_text"):
        parts.append("<resume>\n" + cfg["_resume_text"] + "\n</resume>")
    docs = cfg.get("_documents", [])
    playbooks = [(n, t) for n, t in docs if "playbook" in n.lower()]
    if playbooks:
        parts.append("The interview playbook below comes from the candidate's own past interviews (the ones marked "
                     "as offers were the most successful). Prefer its proven stories, facts and "
                     "phrasings; it is the most reliable picture of how the candidate really talks about their work. "
                     "If it disagrees with the resume on a detail, follow the resume for dates and titles.")
    for name, text in playbooks + [(n, t) for n, t in docs if "playbook" not in n.lower()]:
        parts.append(f'<document name="{name}">\n{text}\n</document>')
    return "\n\n".join(p for p in parts if p)


def resolve_context(cfg: dict) -> dict:
    """Load the texts of the chosen resume and documents into the session config (kept in memory)."""
    cfg = dict(cfg)
    r = get_item("resumes", cfg.get("resume_id") or "") if cfg.get("resume_id") else None
    cfg["_resume_text"] = r["text"][:40000] if r else ""
    docs = []
    ids = cfg.get("document_ids")
    chosen = [it["id"] for it in list_items("documents")] if ids == "all" else (ids if isinstance(ids, list) else [])
    for did in chosen:
        full = get_item("documents", did)
        if full:
            docs.append((full["name"], full["text"][:60000 if "playbook" in full["name"].lower() else 30000]))
    cfg["_documents"] = docs
    return cfg


QUESTION_START = re.compile(
    r"^(what|why|how|when|where|who|which|can|could|would|will|do|does|did|is|are|have|has|tell|describe|explain|"
    r"walk|give|share|talk|imagine|suppose|let'?s|please|"
    r"was|warum|wie|wann|wo|wer|welche|welcher|welches|können|könnten|kannst|würden|würdest|haben|hast|sind|bist|"
    r"erzählen|erzähle|erzähl|erklären|erkläre|erklär|beschreiben|beschreibe|beschreib|stellen|stelle|stell|nennen|nenne|"
    r"sag|sagen|zeig|zeigen|gibt|gib|introduce|"
    r"nə|niyə|necə|nə vaxt|harada|kim|hansı|izah|danış|de)\b", re.I)
QUESTION_ANY = re.compile(
    r"\b(tell me|walk me through|can you|could you|would you|what is|what are|how do|how would|why did|"
    r"what do you|have you|do you|erzählen sie|können sie|könnten sie|was ist|wie würden|haben sie|"
    r"warum haben|was machen|sizcə|izah edin|danışın)\b", re.I)


def _words_of(text: str) -> list[str]:
    return re.findall(r"\w+", text.lower())


def looks_like_question(text: str) -> bool:
    t = text.strip()
    if len(t.split()) < 3:
        return False
    if "?" in t:
        return True
    last = re.split(r"(?<=[.!])\s+", t)[-1]
    return bool(QUESTION_START.match(last) or QUESTION_ANY.search(t[-220:]))


class Copilot:
    """Per recording: watches the interviewer channel and writes answers."""

    WAIT = 0.15  # seconds to wait after a question's final text (the voice detector already saw the pause)
    TURNS = 10   # conversation messages sent with every answer
    HISTORY = 3  # earlier questions and answers sent as memory

    def __init__(self, session, cfg: dict, settings: dict, hub):
        self.session = session
        self.cfg = resolve_context(cfg)
        self.settings = settings
        self.hub = hub
        self.system = build_system(self.cfg)
        self.pending: list[str] = []       # interviewer text since the last answer
        self.pending_at = 0.0
        self.lock = threading.Lock()
        self.stopped = False
        self.shots_dir = session.folder / "shots"
        self.spec: dict | None = None      # answer started from the live caption, waiting for the final text
        threading.Thread(target=self._watch, daemon=True, name="copilot").start()
        threading.Thread(target=self._warm_up, daemon=True, name="copilot-warmup").start()

    def _warm_up(self) -> None:
        """One tiny request at the start: the provider caches the long instructions + resume, so the
        first real answer starts as fast as the later ones (3.1 s -> ~1 s measured)."""
        try:
            stream_chat(self.settings, self.cfg.get("provider"), self.cfg.get("model", ""), self.system,
                        [{"role": "user", "content": "(connection check) Reply with: ok"}], lambda _: None, fast=True)
        except Exception:
            log.info("Copilot warm-up failed", exc_info=True)

    @property
    def mid(self) -> str:
        return self.session.meeting["id"]

    def on_segment(self, seg: dict) -> None:
        if seg["channel"] == "system" and self._settle_spec(seg):
            return
        if seg["channel"] == "system":
            with self.lock:
                self.pending.append(seg["text"])
                self.pending_at = time.monotonic()
                # how long after the words were spoken their final text arrived (diagnostics)
                self.heard_delay = round(self.session.elapsed - seg["end"], 2)
        elif seg["channel"] == "mic" and len(seg["text"].split()) >= 6:
            with self.lock:
                self.pending.clear()             # you already answered; start fresh

    def on_sentence_closed(self, start: float, end: float) -> None:
        """The voice detector just closed an interviewer sentence; its final text needs ~0.7 s more. If the
        live caption already reads like a question, start the answer now (speculative start)."""
        if not self.cfg.get("auto_generate", True) or self.spec:
            return
        st = self.session.agreement.get("system")
        if not st or abs(st.get("start", -1) - start) > 0.6 or not st.get("prev"):
            return
        caption = " ".join(st["prev"])
        if len(caption.split()) < 3:
            return
        with self.lock:
            text = " ".join(self.pending + [caption])
            if not looks_like_question(text):
                return
            before = list(self.pending)
            self.pending.clear()
        item = self.answer(text, auto=True)
        item["speculative"] = True
        item["lag"], item["heard"] = round(self.session.elapsed - end, 2), 0.0   # no wait for the final text
        self.spec = {"item": item, "start": start, "caption": caption, "before": before}

    def _settle_spec(self, seg: dict) -> bool:
        """True when this final segment belonged to the speculative answer (handled here)."""
        spec = self.spec
        if not spec or abs(seg["start"] - spec["start"]) > 0.6:
            return False
        self.spec = None
        item = spec["item"]
        said, heard = _words_of(spec["caption"]), _words_of(seg["text"])
        same = len(set(said) & set(heard)) / max(1, len(set(heard)))
        if same >= 0.8 and len(heard) <= len(said) + 2:
            item["question"] = " ".join(spec["before"] + [seg["text"]]).strip()
            item.pop("speculative", None)
            self.hub.send({"type": "answer_question", "meeting_id": self.mid, "id": item["id"], "question": item["question"]})
            return True
        # the caption missed words: drop that answer and answer the real question
        item["cancelled"] = True
        answers = self.session.meeting.get("answers", [])
        if item in answers:
            answers.remove(item)
        self.hub.send({"type": "answer_removed", "meeting_id": self.mid, "id": item["id"]})
        with self.lock:
            self.pending = spec["before"] + self.pending
        return False

    def _watch(self) -> None:
        while not self.stopped:
            time.sleep(0.15)
            if not self.cfg.get("auto_generate", True):
                continue
            with self.lock:
                if not self.pending or time.monotonic() - self.pending_at < self.WAIT:
                    continue
                waiting = getattr(self.session, "system_speaking", False) or getattr(self.session, "system_inflight", 0)
                if waiting and time.monotonic() - self.pending_at < 4:
                    continue                     # still talking, or the next sentence's text is on its way
                text = " ".join(self.pending)
                if not looks_like_question(text):
                    if len(text.split()) > 120:
                        self.pending = self.pending[-3:]
                    continue
                self.pending.clear()
            self.answer(text, auto=True)

    def recent_turns(self) -> str:
        """The last messages of the conversation, one per speaker turn."""
        turns: list[list] = []
        for s in sorted(self.session.meeting["segments"], key=lambda s: s["start"]):
            who = "Interviewer" if s["channel"] == "system" else "You"
            if turns and turns[-1][0] == who:
                turns[-1][1] += " " + s["text"]
            else:
                turns.append([who, s["text"]])
        return "\n".join(f"{who}: {text}" for who, text in turns[-self.TURNS:])

    def save_shot(self, data: bytes) -> str:
        self.shots_dir.mkdir(exist_ok=True)
        name = f"{int(self.session.elapsed * 1000):09d}.jpg"
        (self.shots_dir / name).write_bytes(data)
        return name

    def _shot_b64(self, name: str) -> str | None:
        import base64

        path = self.shots_dir / name
        return base64.b64encode(path.read_bytes()).decode() if path.exists() else None

    def answer(self, question: str = "", auto: bool = False, image: bytes | None = None) -> dict:
        """Write an answer to `question` (or the interviewer's latest words) and stream it out."""
        if not question.strip():
            segs = sorted(self.session.meeting["segments"], key=lambda s: s["start"])
            last = [s["text"] for s in segs if s["channel"] == "system"][-3:]
            question = " ".join(last) if last and not image else ""
            with self.lock:
                self.pending.clear()
        item = {"id": uuid.uuid4().hex[:8], "question": question.strip(), "text": "", "auto": auto,
                "at": round(self.session.elapsed, 1), "state": "running"}
        heard = [x["end"] for x in self.session.meeting["segments"] if x["channel"] == "system"]
        if auto and heard:
            item["lag"] = round(self.session.elapsed - max(heard), 2)   # question ended -> answer started
            item["heard"] = getattr(self, "heard_delay", None)                # of which: speech-to-text
        if image:
            item["image"] = self.save_shot(image)
        meeting = self.session.meeting
        meeting.setdefault("answers", []).append(item)
        self.hub.send({"type": "answer", "meeting_id": self.mid, "answer": dict(item)})
        threading.Thread(target=self._run, args=(item,), daemon=True, name="copilot-answer").start()
        return item

    def _messages(self, item: dict) -> list[dict]:
        answers = self.session.meeting.get("answers", [])
        earlier = [a for a in answers if a is not item and a.get("state") == "done" and a.get("text")][-self.HISTORY:]
        msgs: list[dict] = []
        for a in earlier:                                     # memory: earlier questions and my answers
            msgs.append({"role": "user", "content": f"(earlier) Question: {a['question'] or '[screenshot]'}"})
            msgs.append({"role": "assistant", "content": a["text"][:1500]})
        text = f"<conversation>\n{self.recent_turns() or '(nothing said yet)'}\n</conversation>\n\n"
        parts: list[dict] = []
        if item.get("image"):
            shots = [a["image"] for a in answers if a.get("image") and a is not item]
            prev = shots[-1] if shots else None
            if prev and (b64 := self._shot_b64(prev)):
                parts += [{"type": "text", "text": "Earlier screenshot (context):"}, {"type": "image", "data": b64}]
            if b64 := self._shot_b64(item["image"]):
                parts += [{"type": "text", "text": "Current screenshot of my screen:"}, {"type": "image", "data": b64}]
            text += ("The question / request: " + item["question"] + "\n" if item["question"] else
                     "Answer what the current screenshot asks for (e.g. solve the coding task), using the conversation as context.\n")
        else:
            text += f"The interviewer's latest question or request:\n<question>\n{item['question']}\n</question>\n"
        return msgs + [{"role": "user", "content": [{"type": "text", "text": text}] + parts if parts else text}]

    def _run(self, item: dict) -> None:
        t0 = time.monotonic()
        buf, last_send = [], 0.0

        def emit(piece: str, flush: bool = False) -> None:
            nonlocal last_send
            if piece:
                buf.append(piece)
                item["text"] += piece
                if "first" not in item:
                    item["first"] = round(time.monotonic() - t0, 2)
            now = time.monotonic()
            if buf and (flush or now - last_send > 0.05):
                self.hub.send({"type": "answer_delta", "meeting_id": self.mid, "id": item["id"], "delta": "".join(buf)})
                buf.clear()
                last_send = now

        provider, model = self.cfg.get("provider"), self.cfg.get("model", "")
        head = {"done": False, "buf": ""}

        def on_piece(piece: str) -> None:
            """The first line is the question title ("# ..."); the rest is the answer."""
            if head["done"]:
                return emit(piece)
            head["buf"] += piece
            if "\n" not in head["buf"] and len(head["buf"]) < 120:
                return
            head["done"] = True
            first, _, rest = head["buf"].partition("\n")
            if first.lstrip().startswith("#"):
                item["title"] = first.strip().lstrip("#").strip()
                self.hub.send({"type": "answer_title", "meeting_id": self.mid, "id": item["id"], "title": item["title"]})
                emit(rest.lstrip("\n"))
            else:
                emit(head["buf"])

        try:
            # One stream without "thinking first": first words after ~1 s, a full answer after ~3 s
            # (measured; ParakeetAI: 1.25 s and ~3 s). Screenshots (code, diagrams) get a little reasoning.
            stream_chat(self.settings, provider, model, self.system, self._messages(item), on_piece,
                        fast=not item.get("image"), cancelled=lambda: item.get("cancelled", False))
            if item.get("cancelled"):
                return                                   # replaced: no "done" (it would bring the card back)
            if not head["done"]:                       # a reply shorter than one line
                head["done"] = True
                emit(head["buf"].strip().lstrip("#").strip())
            emit("", flush=True)
            item["text"] = item["text"].strip()
            item["state"] = "done"
            item["seconds"] = round(time.monotonic() - t0, 1)
        except Exception as exc:
            log.exception("Answer failed")
            item["state"], item["error"] = "error", str(exc)
        if not item.get("cancelled"):
            self.hub.send({"type": "answer_done", "meeting_id": self.mid, "answer": dict(item)})

    def stop(self) -> None:
        self.stopped = True
